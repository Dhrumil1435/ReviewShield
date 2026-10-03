import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_weekly_report_5():
    doc = Document()
    
    # Page Setup - Standard Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Font Styles
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0, 0, 0)

    # 1. Header Section
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_after = Pt(2)
    r_univ = p_univ.add_run("CHAROTAR UNIVERSITY OF SCIENCE &\nTECHNOLOGY")
    r_univ.bold = True
    r_univ.font.size = Pt(14)
    r_univ.font.color.rgb = RGBColor(0, 51, 102)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(12)
    r_inst = p_inst.add_run("DEVANG PATEL INSTITUTE OF ADVANCE TECHNOLOGY\nAND RESEARCH")
    r_inst.bold = True
    r_inst.font.size = Pt(12)
    r_inst.font.color.rgb = RGBColor(0, 51, 102)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(18)
    r_title = p_title.add_run("WEEKLY REPORT - 5")
    r_title.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths: col 0 = 3.25 in, col 1 = 3.25 in
    col_widths = [Inches(3.25), Inches(3.25)]

    table_data = [
        [("Project ID: 32", True), ("Student ID: D25DITT095 , 24DIT070 ,\n24DIT029", True)],
        [("From Date: 11 / 08 / 2026", True), ("To Date: 17 / 08 / 2026", True)],
        [("Semester: 5", True), ("Internship ID:", True)]
    ]

    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            text, is_bold = table_data[r_idx][c_idx]
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.bold = is_bold
            r.font.size = Pt(11)

    # Set Borders for Metadata Table
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
            f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 3. Work Done Heading
    p_work_head = doc.add_paragraph()
    p_work_head.paragraph_format.space_before = Pt(12)
    p_work_head.paragraph_format.space_after = Pt(12)
    r_wh = p_work_head.add_run("Work done from Date: 11/08/2026 to 17/08/2026")
    r_wh.bold = True
    r_wh.font.size = Pt(13)

    # Work Items
    items = [
        ("1. Integrated Meta's RoBERTa Deep Learning Model (src/roberta_engine.py): ",
         "Implemented a standalone contextual sentiment engine using HuggingFace Transformers (cardiffnlp/twitter-roberta-base-sentiment-latest) to compute contextual sentiment probabilities (Negative, Neutral, Positive) and compound sentiment scores (-1.0 to +1.0)."),
        
        ("2. Engineered VADER vs RoBERTa Sentiment Dissonance Features (src/feature_extraction.py): ",
         "Upgraded the feature extraction pipeline to compute 3 new deep learning features per review: roberta_sentiment, roberta_rating_sentiment_gap, and vader_roberta_dissonance (|VADER - RoBERTa|). This isolates deceptive reviews that use surface-level VADER hype words without genuine RoBERTa contextual sentiment."),
        
        ("3. Migrated PostgreSQL Schema & Database Execution (src/db_setup.py): ",
         "Executed auto-migrations on the engineered_features table to append RoBERTa columns. Extracted and stored updated feature vectors across 40,412 review records in PostgreSQL idempotently."),
        
        ("4. Developed Hybrid TF-IDF + Stylometric Model Pipeline (src/train_model.py): ",
         "Built a Scikit-Learn ColumnTransformer pipeline combining StandardScaler (on 9 numerical stylometric/RoBERTa features) with TfidfVectorizer (top 2,000 unigrams + bigrams). Achieved major model performance gains across all benchmarks:"),
        
        ("5. Benchmark Evaluation Metrics: ",
         "Benchmarked classifiers on the hybrid feature set, achieving outstanding performance:\n"
         "• Logistic Regression (Best Model): Accuracy: 89.67% | Precision: 90.48% | Recall: 88.66% | F1-Score: 89.56% | ROC-AUC: 0.9609 (96.09%)\n"
         "• SGDClassifier (LogLoss): Accuracy: 88.44% | Precision: 88.44% | Recall: 88.44% | ROC-AUC: 0.9479\n"
         "• Random Forest Classifier: Accuracy: 82.89% | Precision: 86.97% | Recall: 77.35% | ROC-AUC: 0.9132"),
        
        ("6. Built Interactive Batch CSV Upload Scanner in Streamlit (src/app.py): ",
         "Designed a new '📂 Batch CSV Scanner' page in the web app. Users can upload bulk CSV datasets, dynamically map review text and rating columns, track real-time progress, inspect filterable scan tables, and export results to CSV."),
        
        ("7. Implemented Word-Level Explainable AI (XAI) Scanner (src/inference.py): ",
         "Developed an XAI feature attribution scanner that extracts model TF-IDF weights and renders color-coded HTML text highlights (🔴 Red for deceptive trigger words like 'must buy', 🟢 Green for authentic signal words) inside the Streamlit Scanner UI."),
        
        ("8. Version Control & GitHub Synchronization: ",
         "Committed and pushed the complete hybrid ML pipeline, RoBERTa transformer engine, XAI scanner, and Streamlit dashboard to GitHub repository: https://github.com/Dhrumil1435/ReviewShield.git.")
    ]

    for bold_prefix, text in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.size = Pt(11.5)
        
        r_text = p.add_run(text)
        r_text.font.size = Pt(11.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 4. Plans for Next Week
    p_plan_head = doc.add_paragraph()
    p_plan_head.paragraph_format.space_before = Pt(12)
    p_plan_head.paragraph_format.space_after = Pt(12)
    r_ph = p_plan_head.add_run("Plans for next week: Date: 18/08/2026 to 24/08/2026")
    r_ph.bold = True
    r_ph.font.size = Pt(13)

    plans = [
        ("1. Develop FastAPI REST API Endpoint (src/api.py): ", "Create a lightweight POST /predict REST API route with interactive Swagger documentation (/docs) for third-party e-commerce platform integrations."),
        ("2. Integrate Word Cloud & Radar Chart Analytics: ", "Enhance Streamlit Model Insights page with interactive Word Clouds comparing top deceptive vs authentic keywords and stylometric radar charts."),
        ("3. Complete Comprehensive README & Architectural Documentation: ", "Update repository documentation with system architecture diagrams, updated 89.67% accuracy benchmarking tables, and local installation instructions."),
        ("4. Prepare Final Project Presentation Slide Deck: ", "Draft final demonstration slides, high-resolution UI screenshots, and live project presentation materials for internal review.")
    ]

    for bold_prefix, text in plans:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.size = Pt(11.5)
        
        r_text = p.add_run(text)
        r_text.font.size = Pt(11.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # 5. References Section
    p_ref_head = doc.add_paragraph()
    p_ref_head.paragraph_format.space_after = Pt(8)
    r_ref = p_ref_head.add_run("References:")
    r_ref.bold = True
    r_ref.font.size = Pt(12)

    refs = [
        ("HuggingFace Transformers (RoBERTa): ", "https://huggingface.co/cardiffnlp/twitter-roberta-base-sentiment-latest"),
        ("Scikit-Learn ColumnTransformer & TF-IDF: ", "https://scikit-learn.org/stable/modules/generated/sklearn.compose.ColumnTransformer.html"),
        ("Streamlit Documentation: ", "https://docs.streamlit.io/"),
        ("VADER Sentiment Analysis: ", "https://github.com/cjhutto/vaderSentiment"),
        ("PostgreSQL Documentation: ", "https://www.postgresql.org/docs/"),
        ("ReviewShield GitHub Repository: ", "https://github.com/Dhrumil1435/ReviewShield.git")
    ]

    for title, url in refs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.size = Pt(11)
        
        r_u = p.add_run(url)
        r_u.font.size = Pt(11)
        r_u.font.color.rgb = RGBColor(0, 102, 204)
        r_u.underline = True

    # 6. Signatures Footer
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(40)
    
    r_sig_s = p_sig.add_run("Signature of Student")
    r_sig_s.bold = True
    r_sig_s.font.size = Pt(11.5)
    
    r_space = p_sig.add_run("\t\t\t\t\t\t\t\t")
    
    r_sig_m = p_sig.add_run("Signature of Mentor")
    r_sig_m.bold = True
    r_sig_m.font.size = Pt(11.5)

    # Save Document
    output_path = "Weekly_Report_5_ReviewShield.docx"
    doc.save(output_path)
    print(f"Weekly Report 5 Word document generated successfully at {output_path}")

if __name__ == "__main__":
    generate_weekly_report_5()
